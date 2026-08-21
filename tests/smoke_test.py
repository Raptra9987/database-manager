# -*- coding: utf-8 -*-
"""
Headless checks - these run on GitHub Actions as well.

Run with:  python tests/smoke_test.py
"""
import os
import re
import sys
import json
import shutil
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import database_manager as dm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES = os.path.join(ROOT, "samples")

SANDBOX = tempfile.mkdtemp(prefix="dm_test_")
dm.CONFIG_DIR = os.path.join(SANDBOX, "cfg")
dm.CONFIG_PATH = os.path.join(dm.CONFIG_DIR, "settings.json")
dm.LOG_PATH = os.path.join(dm.CONFIG_DIR, "errors.log")

passed, failed = [], []


def check(name, cond, extra=""):
    (passed if cond else failed).append(f"{name} {extra}".strip())
    print(("  ok    " if cond else "  FAIL  ") + name + (f"   {extra}" if extra else ""))


# --------------------------------------------------------------- translations
print("\n[1] Translations")
broken = [k for k, v in dm.TR.items() if not isinstance(v, tuple) or len(v) != 3]
check("every string has all three languages", not broken, str(broken[:5]))
check("no empty translation",
      not [k for k, v in dm.TR.items() if any(not str(x).strip() for x in v)])

mismatched = [k for k, v in dm.TR.items()
              if len({frozenset(re.findall(r"\{(\w+)\}", s)) for s in v}) > 1]
check("placeholders line up across languages", not mismatched, str(mismatched))
print(f"        {len(dm.TR)} strings x 3 languages = {len(dm.TR) * 3} translations")

for code in dm.LANG_ORDER:
    dm.LANG = code
    check(f"{code}: headers complete", len(dm.headers(dm.default_schema())) == 5)
dm.LANG = "uz"

# ------------------------------------------------------------------ templates
print("\n[2] Templates")
for code, tpl in dm.TEMPLATES.items():
    s = dict(tpl, name=code)
    check(f"{code}: duplicate key is one of the columns",
          s["key_col"] in dm.schema_fields(s), str(s["key_col"]))
    check(f"{code}: has a row-number and a source column",
          dm.role_col(s, "no") and dm.role_col(s, "src"))

# --------------------------------------------------------------- reading files
print("\n[3] Reading the sample files")
contacts = dm.default_schema()
for name, expected in (("existing_database.xlsx", 14), ("contacts_1.xlsx", 3),
                       ("contacts_2_no_header.xlsx", 3), ("contacts_3.txt", 4),
                       ("contacts_4.vcf", 3)):
    rows = dm.load_rows(os.path.join(SAMPLES, name), contacts, with_src=True)
    check(name, len(rows) == expected, f"{len(rows)} / expected {expected}")

rows = dm.load_rows(os.path.join(SAMPLES, "existing_database.xlsx"), contacts,
                    with_src=True)
check("phone column landed correctly", rows[0]["phone"] == "+998901112233",
      rows[0]["phone"])
check("source column landed correctly", rows[0]["source"] == "901234567",
      rows[0]["source"])

# ----------------------------------------------------------- phone formatting
print("\n[4] Phone formatting")
for raw, expected in (("901112233", "+998 90 111 22 33"),
                      ("998901112233", "+998 90 111 22 33"),
                      ("+998 90-111-22-33", "+998 90 111 22 33"),
                      ("79161234567", "+79161234567"),
                      ("12345", "12345"),
                      ("", "")):
    check(f'"{raw}"', dm.pretty_phone(raw) == expected, dm.pretty_phone(raw))
check("applying it twice changes nothing",
      dm.pretty_phone(dm.pretty_phone("901112233")) == "+998 90 111 22 33")

# --------------------------------------------------------------- duplicate key
print("\n[5] Duplicate key")
phone = {"cols": [dm.col("p", 100, "phone")], "key_col": "p", "key_mode": "phone"}
text = {"cols": [dm.col("c", 100)], "key_col": "c", "key_mode": "text"}
number = {"cols": [dm.col("c", 100)], "key_col": "c", "key_mode": "number"}
check("phone: the dialling code does not matter",
      dm.row_key({"p": "+998901112233"}, phone) == dm.row_key({"p": "901112233"}, phone))
check("text: case does not matter",
      dm.row_key({"c": "ABC"}, text) == dm.row_key({"c": "abc"}, text))
check("number: punctuation is dropped", dm.row_key({"c": "SKU-100"}, number) == "100")
check("no key means no duplicate check",
      dm.row_key({"c": "A"}, {"cols": [], "key_col": None}) == "")

# ------------------------------------------------------------ suspicious rows
print("\n[6] Suspicious rows")
for row, expected, label in (
        ({"name": "Ali", "phone": "+998901112233"}, None, "a good row"),
        ({"name": "Ali", "phone": ""}, "phone_yoq", "no phone"),
        ({"name": "Ali", "phone": "12345"}, "qisqa", "too short"),
        ({"name": "Ali", "phone": "998 ABC"}, "harf", "contains letters"),
        ({"name": "", "phone": "+998901112233"}, "ism_yoq", "no name")):
    check(label, dm.row_problem(row, contacts) == expected,
          str(dm.row_problem(row, contacts)))

# ------------------------------------------------- the schema travels with the file
print("\n[7] The schema travels with the file")
from openpyxl import Workbook

path = os.path.join(SANDBOX, "sample.xlsx")
schema = {"name": "custom",
          "cols": [dm.col("no", 55, "no", "h_no"),
                   dm.col("book", 240, title="Book title"),
                   dm.col("isbn", 160, title="ISBN"),
                   dm.col("source", 180, "src", "h_source")],
          "key_col": "isbn", "key_mode": "text"}

wb = Workbook()
ws = wb.active
ws.append(dm.headers(schema))
ws.append([1, "War and Peace", "978-1", "901"])
sheet = wb.create_sheet(dm.SCHEMA_SHEET)
sheet["A1"] = json.dumps(schema, ensure_ascii=False)
sheet.sheet_state = "hidden"
wb.save(path)
wb.close()

restored = dm.read_schema(path)
check("schema restored", restored is not None)
check("columns identical", dm.schema_fields(restored) == dm.schema_fields(schema),
      str(dm.schema_fields(restored)))
check("custom header kept", dm.col_title(restored["cols"][1]) == "Book title")
loaded = dm.load_rows(path, restored, with_src=True)
check("row read back", len(loaded) == 1 and loaded[0]["isbn"] == "978-1", str(loaded))

# a foreign Excel file: columns come from the header row
foreign = os.path.join(SANDBOX, "foreign.xlsx")
wb = Workbook()
ws = wb.active
ws.append(["City", "Population", "Region"])
ws.append(["Tashkent", "2900000", "Tashkent"])
wb.save(foreign)
wb.close()
fs = dm.read_schema(foreign)
check("schema built from a foreign file", fs is not None and len(fs["cols"]) == 3,
      str(fs and [dm.col_title(c) for c in fs["cols"]]))
check("foreign file read", len(dm.load_rows(foreign, fs, with_src=True)) == 1)

# ------------------------------------------------------- settings and error log
print("\n[8] Settings and the error log")
dm.save_config({"lang": "ru", "recent": ["a.xlsx"]})
check("written and read back", dm.load_config().get("lang") == "ru")
with open(dm.CONFIG_PATH, "w", encoding="utf-8-sig") as f:
    f.write(json.dumps({"lang": "en"}))
check("a file written with a BOM still loads", dm.load_config().get("lang") == "en")
with open(dm.CONFIG_PATH, "w", encoding="utf-8") as f:
    f.write("{broken")
check("a corrupt file does not crash the app", dm.load_config() == {})

try:
    raise ValueError("test error")
except ValueError as e:
    dm.log_error(type(e), e, e.__traceback__)
check("the error reached the log",
      "test error" in open(dm.LOG_PATH, encoding="utf-8").read())

# ---------------------------------------------------------------- icon glyphs
print("\n[9] Icon glyphs")
glyphs = [v for k, v in vars(dm).items()
          if k.startswith("IC_") and isinstance(v, str)]
check("each is a single character", all(len(g) == 1 for g in glyphs),
      f"{len(glyphs)} icons")
check("all inside the Segoe icon range",
      all(0xE000 <= ord(g) <= 0xF8FF for g in glyphs))

# ------------------------------------------------------------- missing names
# Every global a function reaches for must actually exist. A rename that
# updates the definition but misses a call site shows up here instead of
# crashing the packaged app on startup.
print("\n[10] Every referenced global exists")
import dis
import builtins
import types

def globals_used(fn):
    try:
        code = fn.__code__
    except AttributeError:
        return set()
    used = {i.argval for i in dis.get_instructions(code)
            if i.opname in ("LOAD_GLOBAL", "STORE_GLOBAL")}
    for const in code.co_consts:          # nested functions and lambdas
        if isinstance(const, types.CodeType):
            used |= {i.argval for i in dis.get_instructions(const)
                     if i.opname == "LOAD_GLOBAL"}
    return used

def ours(obj):
    """Only what this module defines - imported classes carry their own globals."""
    return getattr(obj, "__module__", None) == dm.__name__


checked, missing = 0, {}
targets = [v for v in vars(dm).values()
           if isinstance(v, types.FunctionType) and ours(v)]
for cls in [v for v in vars(dm).values() if isinstance(v, type) and ours(v)]:
    targets += [v for v in vars(cls).values()
                if isinstance(v, (types.FunctionType, staticmethod))]
targets = [getattr(f, "__func__", f) for f in targets]

for fn in targets:
    checked += 1
    for n in globals_used(fn):
        if not hasattr(dm, n) and not hasattr(builtins, n):
            missing.setdefault(fn.__qualname__, []).append(n)

check(f"{checked} functions reference only names that exist", not missing,
      "; ".join(f"{k}: {v}" for k, v in list(missing.items())[:5]))

shutil.rmtree(SANDBOX, ignore_errors=True)
print("\n" + "=" * 52)
print(f"PASSED: {len(passed)}    FAILED: {len(failed)}")
for f in failed:
    print("   - " + f)
sys.exit(1 if failed else 0)
